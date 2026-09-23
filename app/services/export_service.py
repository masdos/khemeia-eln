from __future__ import annotations

import logging
import sqlite3
from collections.abc import Iterator, Sequence
from pathlib import Path
from typing import Any, Protocol

from docx import Document
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from app.repositories import (
    attachment_repository,
    equipment_repository,
    experiment_repository,
    project_repository,
    protocol_repository,
    reagent_repository,
    report_repository,
)

logger = logging.getLogger(__name__)


class ExperimentNotFoundError(ValueError):
    """Raised when exporting an experiment that does not exist."""


class ReportNotFoundError(ValueError):
    """Raised when exporting a report that does not exist."""


class ExperimentRepository(Protocol):
    """Data access contract for experiment reads required by ExportService."""

    def get_by_id(self, experiment_id: int) -> dict[str, Any] | None: ...


class ReagentRepository(Protocol):
    """Data access contract for reagent reads required by ExportService."""

    def get_by_experiment(self, experiment_id: int) -> Sequence[dict[str, Any]]: ...


class EquipmentRepository(Protocol):
    """Data access contract for equipment reads required by ExportService."""

    def get_by_experiment(self, experiment_id: int) -> Sequence[dict[str, Any]]: ...


class ProjectRepository(Protocol):
    """Data access contract for project reads required by ExportService."""

    def get_by_id(self, project_id: int) -> dict[str, Any] | None: ...


class ProtocolRepository(Protocol):
    """Data access contract for protocol reads required by ExportService."""

    def get_by_id(self, protocol_id: int) -> dict[str, Any] | None: ...


class AttachmentRepository(Protocol):
    """Data access contract for attachment reads required by ExportService."""

    def get_by_experiment(self, experiment_id: int) -> Sequence[dict[str, Any]]: ...


class ReportRepository(Protocol):
    """Data access contract for AI report persistence."""

    def create(
        self,
        project_id: int,
        title: str,
        content_markdown: str = "",
    ) -> int: ...

    def link_to_experiments(
        self, report_id: int, experiment_ids: Sequence[int]
    ) -> None: ...

    def get_by_id(self, report_id: int) -> dict[str, Any] | None: ...


class SqliteExperimentRepository:
    """Adapter that backs the experiment read contract with SQLite functions."""

    def __init__(self, connection: sqlite3.Connection) -> None:
        self._connection = connection

    def get_by_id(self, experiment_id: int) -> dict[str, Any] | None:
        return _row_to_dict(
            experiment_repository.get_by_id(self._connection, experiment_id)
        )


class SqliteReagentRepository:
    """Adapter that backs the reagent read contract with SQLite functions."""

    def __init__(self, connection: sqlite3.Connection) -> None:
        self._connection = connection

    def get_by_experiment(self, experiment_id: int) -> Sequence[dict[str, Any]]:
        rows = reagent_repository.get_by_experiment(self._connection, experiment_id)
        return [_row_to_dict(row) for row in rows]


class SqliteEquipmentRepository:
    """Adapter that backs the equipment read contract with SQLite functions."""

    def __init__(self, connection: sqlite3.Connection) -> None:
        self._connection = connection

    def get_by_experiment(self, experiment_id: int) -> Sequence[dict[str, Any]]:
        rows = equipment_repository.get_by_experiment(self._connection, experiment_id)
        return [_row_to_dict(row) for row in rows]


class SqliteProjectRepository:
    """Adapter that backs the project read contract with SQLite functions."""

    def __init__(self, connection: sqlite3.Connection) -> None:
        self._connection = connection

    def get_by_id(self, project_id: int) -> dict[str, Any] | None:
        return _row_to_dict(project_repository.get_by_id(self._connection, project_id))


class SqliteProtocolRepository:
    """Adapter that backs the protocol read contract with SQLite functions."""

    def __init__(self, connection: sqlite3.Connection) -> None:
        self._connection = connection

    def get_by_id(self, protocol_id: int) -> dict[str, Any] | None:
        return _row_to_dict(
            protocol_repository.get_by_id(self._connection, protocol_id)
        )


class SqliteAttachmentRepository:
    """Adapter that backs the attachment read contract with SQLite functions."""

    def __init__(self, connection: sqlite3.Connection) -> None:
        self._connection = connection

    def get_by_experiment(self, experiment_id: int) -> Sequence[dict[str, Any]]:
        rows = attachment_repository.get_by_experiment(self._connection, experiment_id)
        return [_row_to_dict(row) for row in rows]


class SqliteReportRepository:
    """Adapter that backs the report write contract with SQLite functions."""

    def __init__(self, connection: sqlite3.Connection) -> None:
        self._connection = connection

    def create(
        self,
        project_id: int,
        title: str,
        content_markdown: str = "",
    ) -> int:
        return report_repository.create(
            self._connection, project_id, title, content_markdown
        )

    def link_to_experiments(
        self, report_id: int, experiment_ids: Sequence[int]
    ) -> None:
        report_repository.link_to_experiments(
            self._connection, report_id, experiment_ids
        )

    def get_by_id(self, report_id: int) -> dict[str, Any] | None:
        return _row_to_dict(report_repository.get_by_id(self._connection, report_id))


def _row_to_dict(row: sqlite3.Row | None) -> dict[str, Any] | None:
    if row is None:
        return None

    return dict(row)


class ExportService:
    """Generates Markdown, PDF and DOCX exports of experiments and reports."""

    def __init__(
        self,
        base_dir: Path,
        experiment_repo: ExperimentRepository,
        reagent_repo: ReagentRepository,
        equipment_repo: EquipmentRepository,
        project_repo: ProjectRepository,
        protocol_repo: ProtocolRepository,
        attachment_repo: AttachmentRepository,
        user_name: str = "",
        user_email: str = "",
        user_institution: str = "",
        report_repo: ReportRepository | None = None,
    ) -> None:
        self._base_dir = base_dir
        self._experiment_repo = experiment_repo
        self._reagent_repo = reagent_repo
        self._equipment_repo = equipment_repo
        self._project_repo = project_repo
        self._protocol_repo = protocol_repo
        self._attachment_repo = attachment_repo
        self._user_name = user_name
        self._user_email = user_email
        self._user_institution = user_institution
        self._report_repo = report_repo

    def export_experiment_markdown(self, experiment_id: int) -> Path:
        """Export an experiment to Markdown under BASE_DIR/exports/."""
        experiment = self._experiment_repo.get_by_id(experiment_id)
        if experiment is None:
            raise ExperimentNotFoundError(
                f"Experiment with id {experiment_id} does not exist"
            )

        reagents = self._reagent_repo.get_by_experiment(experiment_id)
        equipment = self._equipment_repo.get_by_experiment(experiment_id)
        project = self._project_repo.get_by_id(experiment.get("project_id"))
        protocol = self._protocol_repo.get_by_id(experiment.get("protocol_id"))
        attachments = self._attachment_repo.get_by_experiment(experiment_id)
        content = self._build_markdown(
            experiment, reagents, equipment, project, protocol, attachments
        )

        file_path = self._exports_dir() / f"experiment_{experiment_id}.md"
        file_path.write_text(content, encoding="utf-8")

        logger.info("Experiment exported to markdown experiment_id=%s", experiment_id)
        return file_path

    def export_experiment_pdf(self, experiment_id: int) -> Path:
        """Export an experiment to PDF under BASE_DIR/exports/."""
        experiment = self._experiment_repo.get_by_id(experiment_id)
        if experiment is None:
            raise ExperimentNotFoundError(
                f"Experiment with id {experiment_id} does not exist"
            )

        reagents = self._reagent_repo.get_by_experiment(experiment_id)
        equipment = self._equipment_repo.get_by_experiment(experiment_id)
        project = self._project_repo.get_by_id(experiment.get("project_id"))
        protocol = self._protocol_repo.get_by_id(experiment.get("protocol_id"))
        attachments = self._attachment_repo.get_by_experiment(experiment_id)
        markdown = self._build_markdown(
            experiment, reagents, equipment, project, protocol, attachments
        )

        file_path = self._exports_dir() / f"experiment_{experiment_id}.pdf"
        _write_markdown_pdf(file_path, markdown)

        logger.info("Experiment exported to pdf experiment_id=%s", experiment_id)
        return file_path

    def export_experiment_docx(self, experiment_id: int) -> Path:
        """Export an experiment to DOCX under BASE_DIR/exports/."""
        experiment = self._experiment_repo.get_by_id(experiment_id)
        if experiment is None:
            raise ExperimentNotFoundError(
                f"Experiment with id {experiment_id} does not exist"
            )

        reagents = self._reagent_repo.get_by_experiment(experiment_id)
        equipment = self._equipment_repo.get_by_experiment(experiment_id)
        project = self._project_repo.get_by_id(experiment.get("project_id"))
        protocol = self._protocol_repo.get_by_id(experiment.get("protocol_id"))
        attachments = self._attachment_repo.get_by_experiment(experiment_id)
        markdown = self._build_markdown(
            experiment, reagents, equipment, project, protocol, attachments
        )

        file_path = self._exports_dir() / f"experiment_{experiment_id}.docx"
        _write_markdown_docx(file_path, markdown)

        logger.info("Experiment exported to docx experiment_id=%s", experiment_id)
        return file_path

    def save_report(
        self,
        markdown_content: str,
        experiment_ids: Sequence[int],
        project_id: int,
        title: str,
    ) -> int:
        """Store an AI report draft in the database and return its identifier.

        Only the database is touched: no file is written. The full content
        is stored so the report can be recovered later.
        """
        self._validate_ai_report_input(markdown_content, experiment_ids)
        if not title or not title.strip():
            raise ValueError("title must not be empty")
        if self._project_repo.get_by_id(project_id) is None:
            raise ValueError(f"Project with id {project_id} does not exist")
        report_repo = self._require_report_repo()

        report_id = report_repo.create(project_id, title.strip(), markdown_content)
        report_repo.link_to_experiments(report_id, list(experiment_ids))

        logger.info("AI report saved report_id=%s project_id=%s", report_id, project_id)
        return report_id

    def export_ai_report_markdown(self, markdown_content: str) -> Path:
        """Export AI generated Markdown to a file under BASE_DIR/exports/."""
        if not markdown_content or not markdown_content.strip():
            raise ValueError("markdown_content must not be empty")

        _, stored_name = _build_ai_report_names("md")
        file_path = self._exports_dir() / stored_name
        content = self._with_user_header(markdown_content)
        file_path.write_text(content, encoding="utf-8")

        logger.info("AI report exported to markdown stored_name=%s", stored_name)
        return file_path

    def export_ai_report_pdf(self, markdown_content: str) -> Path:
        """Export AI generated Markdown as a PDF file under BASE_DIR/exports/."""
        if not markdown_content or not markdown_content.strip():
            raise ValueError("markdown_content must not be empty")

        _, stored_name = _build_ai_report_names("pdf")
        file_path = self._exports_dir() / stored_name
        _write_markdown_pdf(file_path, self._with_user_header(markdown_content))

        logger.info("AI report exported to pdf stored_name=%s", stored_name)
        return file_path

    def export_ai_report_docx(self, markdown_content: str) -> Path:
        """Export AI generated Markdown as a DOCX file under BASE_DIR/exports/."""
        if not markdown_content or not markdown_content.strip():
            raise ValueError("markdown_content must not be empty")

        _, stored_name = _build_ai_report_names("docx")
        file_path = self._exports_dir() / stored_name
        _write_markdown_docx(file_path, self._with_user_header(markdown_content))

        logger.info("AI report exported to docx stored_name=%s", stored_name)
        return file_path

    def export_report_markdown(self, report_id: int) -> Path:
        """Export a saved report to Markdown under BASE_DIR/exports/."""
        report = self._require_report(report_id)
        content = self._with_user_header(report.get("content_markdown") or "")

        file_path = self._exports_dir() / f"report_{report_id}.md"
        file_path.write_text(content, encoding="utf-8")

        logger.info("Report exported to markdown report_id=%s", report_id)
        return file_path

    def export_report_pdf(self, report_id: int) -> Path:
        """Export a saved report to PDF under BASE_DIR/exports/."""
        report = self._require_report(report_id)
        content = self._with_user_header(report.get("content_markdown") or "")

        file_path = self._exports_dir() / f"report_{report_id}.pdf"
        _write_markdown_pdf(file_path, content)

        logger.info("Report exported to pdf report_id=%s", report_id)
        return file_path

    def export_report_docx(self, report_id: int) -> Path:
        """Export a saved report to DOCX under BASE_DIR/exports/."""
        report = self._require_report(report_id)
        content = self._with_user_header(report.get("content_markdown") or "")

        file_path = self._exports_dir() / f"report_{report_id}.docx"
        _write_markdown_docx(file_path, content)

        logger.info("Report exported to docx report_id=%s", report_id)
        return file_path

    def _require_report(self, report_id: int) -> dict[str, Any]:
        report_repo = self._require_report_repo()
        report = report_repo.get_by_id(report_id)
        if report is None:
            raise ReportNotFoundError(f"Report with id {report_id} does not exist")

        return report

    def _user_header_lines(self) -> list[str]:
        from datetime import date

        header = [
            date.today().strftime("%Y-%m-%d"),
            self._user_name,
            self._user_email,
        ]
        if self._user_institution.strip():
            header.append(self._user_institution.strip())
        return header

    def _with_user_header(self, markdown_content: str) -> str:
        return "\n".join([*self._user_header_lines(), "", markdown_content]) + "\n"

    def _require_report_repo(self) -> ReportRepository:
        if self._report_repo is None:
            raise RuntimeError("Report repository is not configured")
        return self._report_repo

    @staticmethod
    def _validate_ai_report_input(
        markdown_content: str, experiment_ids: Sequence[int]
    ) -> None:
        if not markdown_content or not markdown_content.strip():
            raise ValueError("markdown_content must not be empty")
        if not experiment_ids:
            raise ValueError("experiment_ids must contain at least one id")

    def _build_markdown(
        self,
        experiment: dict[str, Any],
        reagents: Sequence[dict[str, Any]],
        equipment: Sequence[dict[str, Any]],
        project: dict[str, Any] | None,
        protocol: dict[str, Any] | None,
        attachments: Sequence[dict[str, Any]],
    ) -> str:
        project_name = project["name"] if project else "_None_"
        protocol_name = protocol["name"] if protocol else "_None_"

        lines = [
            *self._user_header_lines(),
            "",
            f"# {experiment['title']}",
            "",
            f"**Project:** {project_name}",
            f"**Protocol:** {protocol_name}",
            f"**State:** {experiment['state']}",
            "",
            "## Question",
            "",
            experiment.get("question") or "_No question recorded._",
            "",
            "## Experimental Procedure",
            "",
            experiment.get("experimental_procedure_markdown")
            or "_No experimental procedure recorded._",
            "",
            "## Result",
            "",
            experiment.get("result_markdown") or "_No result recorded._",
            "",
            "## Conclusions",
            "",
            experiment.get("conclusions") or "_No conclusions recorded._",
            "",
            "## Reagents",
            "",
        ]
        if reagents:
            lines.extend(self._format_reagent_line(row) for row in reagents)
        else:
            lines.append("_No reagents recorded._")

        lines.extend(["", "## Equipment", ""])
        if equipment:
            lines.extend(f"- {row['name']}" for row in equipment)
        else:
            lines.append("_No equipment recorded._")

        lines.extend(["", "## Attachments", ""])
        if attachments:
            lines.extend(f"- {att['file_name']}" for att in attachments)
        else:
            lines.append("_No attachments._")

        return "\n".join(lines) + "\n"

    def _format_reagent_line(self, row: dict[str, Any]) -> str:
        parts = [f"- {row['name']}"]
        details: list[str] = []
        if row.get("amount_used") is not None:
            details.append(f"{row['amount_used']} {row['unit']}")
        if row.get("lot_number"):
            details.append(f"Lot: {row['lot_number']}")
        if details:
            parts.append(f" ({', '.join(details)})")
        return "".join(parts)

    def _exports_dir(self) -> Path:
        exports_dir = self._base_dir / "exports"
        exports_dir.mkdir(parents=True, exist_ok=True)
        return exports_dir


def _build_ai_report_names(extension: str) -> tuple[str, str]:
    """Build human readable and unique storage names for an AI report."""
    import uuid
    from datetime import datetime, timezone

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    file_name = f"ai_report_{timestamp}.{extension}"
    stored_name = f"{uuid.uuid4().hex}.{extension}"
    return file_name, stored_name


def _split_table_row(line: str) -> list[str]:
    """Split a GitHub Markdown table row into stripped cell values."""
    placeholder = "\0"
    protected = line.replace("\\|", placeholder)
    stripped = protected.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip().replace(placeholder, "|") for cell in stripped.split("|")]


def _is_table_separator(line: str) -> bool:
    """Check whether a line is a GitHub Markdown table separator row."""
    import re

    cells = _split_table_row(line)
    return len(cells) > 0 and all(
        re.fullmatch(r":?-{1,}:?", cell) is not None for cell in cells
    )


def _collect_table_block(
    lines: Sequence[str], start: int
) -> tuple[list[str], list[list[str]], int] | None:
    """Collect a GitHub Markdown table starting at the given line index.

    Return the header cells, body rows and number of consumed lines,
    or None when the lines do not form a table.
    """
    first = lines[start].strip()
    if not first.startswith("|") or "|" not in first[1:]:
        return None
    if start + 1 >= len(lines) or not _is_table_separator(lines[start + 1]):
        return None

    header = _split_table_row(first)
    rows: list[list[str]] = []
    consumed = 2
    for raw_line in lines[start + 2 :]:
        stripped = raw_line.strip()
        if not stripped.startswith("|") or "|" not in stripped[1:]:
            break
        rows.append(_split_table_row(stripped))
        consumed += 1

    width = len(header)
    normalized = [
        (row + [""] * width)[:width] if len(row) != width else row for row in rows
    ]
    return header, normalized, consumed


def _write_markdown_pdf(file_path: Path, markdown: str) -> None:
    """Render simple Markdown content into a PDF file."""
    styles = getSampleStyleSheet()
    heading_styles = {
        1: ParagraphStyle(
            "Heading1", parent=styles["Heading1"], textColor=colors.HexColor("#1F2937")
        ),
        2: ParagraphStyle(
            "Heading2", parent=styles["Heading2"], textColor=colors.HexColor("#374151")
        ),
        3: ParagraphStyle(
            "Heading3", parent=styles["Heading3"], textColor=colors.HexColor("#4B5563")
        ),
    }
    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        alignment=TA_LEFT,
        spaceAfter=6,
    )

    document = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
        rightMargin=2 * cm,
        leftMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=file_path.stem,
    )

    flowables: list = []
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            flowables.append(Spacer(1, 6))
            index += 1
            continue

        table_block = _collect_table_block(lines, index)
        if table_block is not None:
            header, rows, consumed = table_block
            flowables.append(_build_pdf_table(header, rows, body_style))
            index += consumed
            continue

        if line.startswith("### "):
            flowables.append(
                Paragraph(
                    _convert_inline_markdown(_escape_html(line[4:])), heading_styles[3]
                )
            )
        elif line.startswith("## "):
            flowables.append(
                Paragraph(
                    _convert_inline_markdown(_escape_html(line[3:])), heading_styles[2]
                )
            )
        elif line.startswith("# "):
            flowables.append(
                Paragraph(
                    _convert_inline_markdown(_escape_html(line[2:])), heading_styles[1]
                )
            )
        elif line.startswith("- "):
            flowables.append(
                ListFlowable(
                    [
                        ListItem(
                            Paragraph(
                                _convert_inline_markdown(_escape_html(line[2:])),
                                body_style,
                            ),
                            leftIndent=18,
                        )
                    ],
                    bulletType="bullet",
                    start="•",
                )
            )
        else:
            flowables.append(
                Paragraph(_convert_inline_markdown(_escape_html(line)), body_style)
            )
        index += 1

    document.build(flowables)


def _build_pdf_table(
    header: list[str], rows: list[list[str]], body_style: ParagraphStyle
) -> Table:
    """Build a styled reportlab Table from header and body rows."""
    data = [
        [
            Paragraph(_convert_inline_markdown(_escape_html(cell)), body_style)
            for cell in header
        ]
    ]
    data.extend(
        [
            Paragraph(_convert_inline_markdown(_escape_html(cell)), body_style)
            for cell in row
        ]
        for row in rows
    )
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E7EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F2937")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9CA3AF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ]
        )
    )
    return table


def _write_markdown_docx(file_path: Path, markdown: str) -> None:
    """Render simple Markdown content into a DOCX file."""
    document = Document()
    document.core_properties.title = file_path.stem

    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        if not line.strip():
            index += 1
            continue

        table_block = _collect_table_block(lines, index)
        if table_block is not None:
            header, rows, consumed = table_block
            _add_docx_table(document, header, rows)
            index += consumed
            continue

        if line.startswith("### "):
            _add_docx_paragraph(document, line[4:], style="Heading 3")
        elif line.startswith("## "):
            _add_docx_paragraph(document, line[3:], style="Heading 2")
        elif line.startswith("# "):
            _add_docx_paragraph(document, line[2:], style="Heading 1")
        elif line.startswith("- "):
            _add_docx_paragraph(document, line[2:], style="List Bullet")
        else:
            _add_docx_paragraph(document, line)
        index += 1

    document.save(str(file_path))


def _add_docx_table(
    document: Document, header: list[str], rows: list[list[str]]
) -> None:
    """Add a gridded table with a bold header row to the document."""
    table = document.add_table(rows=1 + len(rows), cols=len(header))
    table.style = "Table Grid"
    for col, cell_text in enumerate(header):
        _set_docx_cell_text(table.cell(0, col), cell_text, bold=True)
    for row_idx, row in enumerate(rows, start=1):
        for col, cell_text in enumerate(row):
            _set_docx_cell_text(table.cell(row_idx, col), cell_text)


def _set_docx_cell_text(cell: Any, text: str, bold: bool = False) -> None:
    """Fill a table cell with **bold** and _italic_ markdown as styled runs."""
    paragraph = cell.paragraphs[0]
    for chunk, chunk_bold, italic in _iter_inline_runs(text):
        run = paragraph.add_run(chunk)
        run.bold = bold or chunk_bold
        run.italic = italic


def _add_docx_paragraph(
    document: Document, text: str, style: str | None = None
) -> None:
    """Add a paragraph with **bold** and _italic_ markdown as styled runs."""
    paragraph = (
        document.add_paragraph(style=style) if style else document.add_paragraph()
    )
    for chunk, bold, italic in _iter_inline_runs(text):
        run = paragraph.add_run(chunk)
        run.bold = bold
        run.italic = italic


def _iter_inline_runs(text: str) -> Iterator[tuple[str, bool, bool]]:
    """Split **bold** and _italic_ markdown into (text, bold, italic) runs."""
    import re

    pattern = re.compile(r"\*\*(.+?)\*\*|(?<!\w)_(.+?)_(?!\w)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            yield text[pos : match.start()], False, False
        if match.group(1) is not None:
            yield match.group(1), True, False
        else:
            yield match.group(2), False, True
        pos = match.end()
    if pos < len(text):
        yield text[pos:], False, False


def _escape_html(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _convert_inline_markdown(text: str) -> str:
    """Convert **bold** and _italic_ markdown to ReportLab HTML tags."""
    import re

    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"(?<!\w)_(.+?)_(?!\w)", r"<i>\1</i>", text)
    return text
